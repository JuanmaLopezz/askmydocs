import pytest
from pathlib import Path
from langchain_core.documents import Document
from app.ingestion.loaders import load_document, _load_txt
from app.ingestion.chunker import chunk_documents


def test_load_txt(tmp_path):
    f = tmp_path / "sample.txt"
    f.write_text("Hello world. This is a test document.", encoding="utf-8")
    docs = load_document(f)
    assert len(docs) == 1
    assert "Hello world" in docs[0].page_content


def test_load_unsupported_extension(tmp_path):
    f = tmp_path / "sample.csv"
    f.write_text("a,b,c", encoding="utf-8")
    with pytest.raises(ValueError, match="not supported"):
        load_document(f)


def test_chunk_documents():
    long_text = "word " * 500
    docs = [Document(page_content=long_text, metadata={"source": "test.txt"})]
    chunks = chunk_documents(docs)
    assert len(chunks) > 1
    for chunk in chunks:
        assert len(chunk.page_content) <= 1200


def test_chunk_preserves_metadata():
    docs = [Document(page_content="word " * 300, metadata={"source": "test.txt", "document_id": "abc"})]
    chunks = chunk_documents(docs)
    for chunk in chunks:
        assert chunk.metadata["document_id"] == "abc"


def test_load_xlsx(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["Name", "Age"])
    ws.append(["Alice", 30])
    ws.append(["Bob", 25])
    path = tmp_path / "sample.xlsx"
    wb.save(path)
    docs = load_document(path)
    assert len(docs) >= 1
    assert "Alice" in docs[0].page_content
